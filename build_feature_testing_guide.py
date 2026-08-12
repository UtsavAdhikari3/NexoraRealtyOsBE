"""Build the Nexora RealtyOS Feature Testing Guide from the QA catalog."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import qa_master_catalog as qa


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "outputs" / "master_qa_documentation"
OUTPUT = OUTPUT_DIR / "Feature_Testing_Guide.docx"

NAVY = "17324D"
BLUE = "246B8E"
CYAN = "D9EEF5"
PALE = "EEF5F8"
INK = "17212B"
MUTED = "5D6B78"
LINE = "CCD8E0"
AMBER = "FCE8CC"
RED = "F8D7DA"
GREEN = "DFF1E3"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=LINE, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_font(run, name="Aptos", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(9.2)
        style.paragraph_format.space_after = Pt(2)

    h1 = styles["Heading 1"]
    h1.font.name = "Aptos Display"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(NAVY)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Aptos Display"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Aptos"
    h3._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    h3._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    h3.font.size = Pt(10.5)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(INK)
    h3.paragraph_format.space_before = Pt(5)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.keep_with_next = True


def configure_section(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("NEXORA REALTYOS  /  MASTER QA GUIDE")
    set_font(r, size=8, color=MUTED, bold=True)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.75))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.3)
    table.columns[1].width = Inches(1.45)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = left.add_run("Current implementation baseline  |  Generated 2026-08-09")
    set_font(r, size=7.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = right.add_run("Page ")
    set_font(r, size=7.5, color=MUTED)
    add_page_field(right)


def add_para(doc, text="", bold=False, italic=False, color=INK, size=None, align=None, before=0, after=5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if level:
            p.paragraph_format.left_indent = Inches(0.25 * level)
        p.add_run(str(item))


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(str(item))


def add_callout(doc, title, text, fill=PALE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper())
    set_font(r, size=8, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(text)
    set_font(r, size=9.2, color=INK)
    set_table_borders(table, accent, "8")
    add_para(doc, after=2)


def add_two_col_table(doc, rows, widths=(1.6, 5.05), header=None):
    table = doc.add_table(rows=1 if header else 0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(widths[0])
    table.columns[1].width = Inches(widths[1])
    if header:
        cells = table.rows[0].cells
        for i, text in enumerate(header):
            set_cell_shading(cells[i], NAVY)
            p = cells[i].paragraphs[0]
            r = p.add_run(text)
            set_font(r, size=8.5, color=WHITE, bold=True)
        set_repeat_table_header(table.rows[0])
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        set_cell_shading(cells[0], PALE)
        if idx % 2:
            set_cell_shading(cells[1], "F8FAFB")
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
        p = cells[0].paragraphs[0]
        r = p.add_run(str(label))
        set_font(r, size=8.4, color=NAVY, bold=True)
        p = cells[1].paragraphs[0]
        r = p.add_run(str(value))
        set_font(r, size=8.5, color=INK)
    set_table_borders(table)
    return table


def add_flow(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7FAFC")
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(lines):
        if i:
            r = p.add_run("\n  |\n  v\n")
            set_font(r, name="Consolas", size=8.2, color=MUTED)
        r = p.add_run(line)
        set_font(r, name="Consolas", size=8.6, color=NAVY, bold=True)
    set_table_borders(table, LINE)


def cover_page(doc, metrics):
    add_para(doc, "MASTER QA DOCUMENTATION", bold=True, color=BLUE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, before=54, after=18)
    add_para(doc, "Feature Testing Guide", bold=True, color=NAVY, size=30, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "Nexora RealtyOS", bold=True, color=INK, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    add_para(doc, "Manual execution handbook for the Django API, agency dashboard, public storefront, scheduler and external integrations", color=MUTED, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=34)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    values = [
        (str(metrics["feature_areas"]), "Feature areas"),
        (str(metrics["total_cases"]), "Workbook cases"),
        (str(metrics["p0"]), "P0 cases"),
        (str(metrics["e2e"]), "E2E journeys"),
    ]
    for i, (value, label) in enumerate(values):
        cell = table.cell(0, i)
        cell.width = Inches(1.65)
        set_cell_shading(cell, PALE if i % 2 == 0 else CYAN)
        set_cell_margins(cell, top=180, start=90, bottom=180, end=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_font(r, size=19, color=NAVY, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(label)
        set_font(r, size=8, color=MUTED, bold=True)
    set_table_borders(table, LINE)

    add_para(doc, "", after=28)
    add_callout(
        doc,
        "Scope boundary",
        "Stripe tests in this guide cover Nexora's agency SaaS subscription only. Property-sale payments and rent collection are not current product workflows. Property owners do not self-publish listings; the agency creates and manages property inventory.",
        fill=AMBER,
        accent="8A4B08",
    )
    add_para(doc, "Version 1.0  |  Current implementation baseline  |  09 August 2026", color=MUTED, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, before=28)
    doc.add_page_break()


def front_matter(doc, metrics):
    doc.add_heading("How to Use This Guide", level=1)
    add_para(doc, "This guide teaches a tester how to execute Nexora RealtyOS manually. The companion workbook is the case repository and execution tracker. This document focuses on setup, navigation, observation points, cross-feature effects, failure simulation and regression impact.")
    add_callout(doc, "Pass/fail discipline", "A test passes only when the UI, API response, database impact and downstream side effects all match. A correct screen with a wrong tenant relationship, duplicate notification or stale public listing is a failure.")

    doc.add_heading("Recommended Test Environment", level=2)
    add_bullets(doc, [
        "Run PostgreSQL 16, Django/Gunicorn API, scheduler, React/Vite agency UI and Next.js storefront through the repository Docker setup.",
        "Use separate test-only SMTP and Meta/Stripe test or mocked integrations. Never publish to production social accounts during routine QA.",
        "Keep browser DevTools Network/Console open. Preserve request/response payloads for every failure.",
        "Use an API client for direct permission, malformed payload, concurrency and retry tests.",
        "Use read-only database access for verification of row counts, tenant IDs, histories, timestamps, unique constraints and rollback.",
        "Capture API/scheduler logs and provider mock call counts for background and integration tests.",
        "Set system/application timezone expectations to Asia/Kathmandu and store UTC-aware timestamps consistently.",
    ])

    doc.add_heading("Reusable Test Data Matrix", level=2)
    rows = [
        ("Agency A", "Everest Homes Nepal; active/paid; website published; slug everest-homes; Nepal localization enabled."),
        ("Agency B", "Himalayan Realty; separate active/paid tenant; unique users, properties, Meta account and customers."),
        ("Agency C", "Suspended/expired subscription; website record retained but public/auth access blocked."),
        ("Super Admin", "platform.admin@nexora.test; role super_admin; no ordinary tenant assumption."),
        ("Owner A / B", "owner.a@nexora.test and owner.b@nexora.test; full tenant administration."),
        ("Manager A", "manager.a@nexora.test; rules, verification, republish, reviews and team management."),
        ("Agent A1 / A2", "Distinct capacities/assignments; A1 public profile; A2 used for unauthorized-access comparisons."),
        ("Agent B1", "Agency B identity used in cross-tenant relationship/IDOR tests."),
        ("Customer A1 / A2", "Separate Agency A portal accounts; saved property/search and appointment fixtures."),
        ("Customer B1", "Agency B token used to prove customer tenant isolation."),
        ("Property A-PUB", "Fresh, verified, available, published listing with Nepal fields and media."),
        ("Property A-PRIVATE", "Draft/expired/unpublished listing with unique facet values to detect public leakage."),
        ("Property B-PUB", "Agency B public listing used in cross-tenant ID substitution."),
        ("Lead A1 / A2", "Assigned to different agents, with source attribution, follow-up and automation timestamps."),
    ]
    add_two_col_table(doc, rows, header=("Fixture", "Reusable Definition"))

    doc.add_heading("Execution Evidence Checklist", level=2)
    add_bullets(doc, [
        "Screenshot or screen recording of visible behavior.",
        "HTTP method, actual URL, request body/headers, status code and response body.",
        "Before/after database counts and key fields for affected models.",
        "Application/scheduler log excerpt with correlation time/record ID.",
        "Email/notification/provider mock evidence and call count.",
        "Tenant A/B comparison for every sensitive workflow.",
        "Actual Result, Status and Bug ID updated in System_Test_Cases.xlsx.",
    ])

    doc.add_heading("System Entry Points and Roles", level=2)
    add_two_col_table(doc, [
        ("Agency dashboard", "React/Vite routes such as /dashboard, /properties, /leads, /inbox, /site-visits, /deals, /appointments, /social-media, /settings and /platform-admin."),
        ("Public storefront", "Next.js and React public agency/property/agent/customer-portal routes backed by /api/public/ endpoints."),
        ("Backend API", "Django REST endpoints under /api/auth, /api/agencies, /api/properties, /api/leads, /api/site-visits, /api/operations, /api/inbox and /api/social-posts."),
        ("Scheduler", "Management commands executed by the Docker scheduler loop every configured interval."),
        ("Roles", "Super Admin; Agency Owner; Agency Manager; Agent; anonymous visitor; agency-scoped CustomerProfile; external Meta/Stripe systems."),
    ], header=("Surface", "Testing Use"))

    doc.add_page_break()
    doc.add_heading("Feature Inventory", level=2)
    inventory = [(f"{i:02d}. {f['name']}", f["purpose"]) for i, f in enumerate(qa.FEATURES, 1)]
    add_two_col_table(doc, inventory, widths=(2.15, 4.5), header=("Feature", "Scope"))
    doc.add_page_break()


def feature_chapter(doc, number, feature):
    if doc.paragraphs and doc.paragraphs[-1].text:
        doc.add_page_break()
    doc.add_heading(f"{number}. {feature['name']}", level=1)

    doc.add_heading("Purpose", level=2)
    add_para(doc, feature["purpose"])

    doc.add_heading("Current Implementation Trace", level=2)
    add_two_col_table(doc, [
        ("Users / Roles", feature["roles"]),
        ("Where to Start", feature["ui"]),
        ("Actual API Routes", feature["endpoints"]),
        ("Primary Models", feature["models"]),
        ("Dependencies", feature["dependencies"]),
    ], header=("Layer", "What the Tester Should Trace"))

    doc.add_heading("Preconditions", level=2)
    add_bullets(doc, [x.strip() for x in feature["prerequisites"].split(";") if x.strip()])

    doc.add_heading("Test Data Needed", level=2)
    add_para(doc, feature["data"])

    doc.add_heading("Normal Feature Flow", level=2)
    flow_lines = [
        feature["roles"].split(";")[0],
        feature["ui"].split(",")[0],
        "Frontend validation and authenticated/public request",
        feature["endpoints"].split(";")[0],
        "Backend permission, tenant and serializer validation",
        feature["models"].split(",")[0],
        "UI refresh plus downstream verification",
    ]
    add_flow(doc, flow_lines)

    doc.add_heading("Detailed Manual Testing Steps", level=2)
    steps = [
        f"Start the documented Docker services and confirm /api/health/ before testing {feature['name']}.",
        f"Seed the required fixtures: {feature['prerequisites']}",
        f"Login or enter the public flow as: {feature['roles']}.",
        f"Navigate to {feature['ui']} and capture the initial list/detail state.",
        f"Perform the normal workflow: {feature['primary_action']}",
        "Keep DevTools Network open and save the exact request, response, status code and timing.",
        f"Verify model changes in {feature['models']} using read-only queries or Django shell inspection.",
        f"Verify expected result: {feature['primary_result']}",
        f"Inspect downstream modules: {feature['dependencies']}.",
        "Refresh the browser and repeat the GET request to prove persistence and remove optimistic UI ambiguity.",
        "Repeat the action as Agent, Manager/Owner, unauthenticated caller and Agency B user where applicable.",
        "Record Actual Result, Status, evidence and Bug ID in the companion workbook.",
    ]
    add_numbered(doc, steps)

    doc.add_heading("Expected Result", level=2)
    add_callout(doc, "Pass condition", feature["primary_result"], fill=GREEN, accent="235C34")

    doc.add_heading("What Else Should Change?", level=2)
    add_bullets(doc, [
        feature["db_impact"],
        feature["side_effects"],
        f"Related feature state should remain consistent across {feature['dependencies']}.",
    ])

    doc.add_heading("What Should NOT Change?", level=2)
    add_bullets(doc, [
        "Agency B records, counts, files, notifications and public output.",
        "Unrelated Agency A records not referenced by the request.",
        "Protected role, payment or ownership fields unless this feature explicitly owns them.",
        "External systems when local validation or permission fails.",
        "Property transaction or rent payment state; those workflows are not implemented unless this is SaaS subscription billing.",
    ])

    doc.add_heading("Negative Testing", level=2)
    add_para(doc, feature["validation"])
    add_bullets(doc, [
        "Omit each required value one at a time, then submit all required values missing.",
        "Use an existing but unauthorized ID rather than only a random missing ID.",
        "Send malformed JSON, invalid choice values and whitespace-only strings directly to the API.",
        "Expire/deactivate authentication or tenant entitlement between page load and submit.",
        "Simulate a 500/timeout and confirm the UI never reports false success.",
    ])

    doc.add_heading("Boundary and Duplicate Testing", level=2)
    add_bullets(doc, [
        f"Current duplicate/idempotency rule: {feature['duplicate_rule']}",
        "For model max_length and Decimal fields, test limit-1, limit and limit+1 using the actual model/serializer constraints.",
        "Test zero, negative, null, empty, whitespace, Unicode Nepali and unusually long values where semantically relevant.",
        "For dates/times, test just before, exactly at and just after the threshold in Asia/Kathmandu.",
        "Send the same request twice sequentially and concurrently; inspect both row and side-effect counts.",
    ])

    doc.add_heading("Permission and Tenant Testing", level=2)
    add_para(doc, feature["permission_rule"])
    add_bullets(doc, [
        "Repeat direct API calls even if the UI hides the control.",
        "Replace the target and every related ID with Agency B fixtures.",
        "Check list, detail, mutation, custom actions, files/downloads and filter-option endpoints separately.",
        "Verify error text does not reveal another tenant's record title, filename or relationship.",
    ])

    doc.add_heading("Dependency Testing", level=2)
    add_para(doc, f"This feature depends on or feeds: {feature['dependencies']}. After a successful mutation, open those modules and verify the intended count, status, attribution, notification, visibility or relationship change. After a failed mutation, verify all remain unchanged.")

    doc.add_heading("Database Verification", level=2)
    add_para(doc, feature["db_impact"])
    add_bullets(doc, [
        "Verify agency_id on every created row and related object.",
        "Check unique constraints and duplicate counts after double submission.",
        "Check history/audit/timestamp rows only for accepted mutations.",
        "For deletion, validate CASCADE, PROTECT and SET_NULL effects and transaction rollback.",
    ])

    doc.add_heading("API Verification", level=2)
    add_para(doc, feature["endpoints"])
    add_bullets(doc, [
        "2xx response body matches the persisted record after re-query.",
        "400 is used for invalid payload/relationship/state; 401 for missing/invalid auth; 403 for role denial; scoped 404 hides inaccessible records.",
        "Pagination/filter/search and response fields remain tenant safe.",
        "No endpoint, method or behavior should be invented if it is absent from these routes.",
    ])

    doc.add_heading("Failure and Recovery Testing", level=2)
    add_bullets(doc, [
        "Disable network after submit and determine whether the server committed before retrying.",
        "Expire access JWT and verify one refresh/retry; fail refresh and verify safe logout.",
        "Force database/provider/email failure at the side-effect boundary and inspect transaction/result consistency.",
        "Run scheduler/provider retry twice and prove logical idempotency where implemented.",
        "If behavior is not determinable, record: Requires clarification / not determined from current implementation.",
    ])

    doc.add_heading("Regression Impact", level=2)
    add_para(doc, f"If {feature['name']} changes, retest its primary workflow, validation, roles, tenant isolation, duplicate/delete behavior, public/private exposure and every dependency listed here: {feature['dependencies']}.")

    special = feature["specials"][0]
    doc.add_heading("Reusable Steps to Reproduce", level=2)
    add_two_col_table(doc, [
        ("Feature", feature["name"]),
        ("Scenario", special[0]),
        ("Precondition", feature["prerequisites"]),
        ("Steps", f"1. Capture initial state. 2. Use role {special[3]}. 3. {special[4]} 4. Inspect API/UI/database/side effects. 5. Refresh and repeat once."),
        ("Expected", special[5]),
        ("Actual", "To be filled by tester."),
        ("Evidence", "Request/response, screenshots, database diff, logs and external mock call count."),
    ], header=("Field", "Reproduction Template"))


def cross_feature_section(doc):
    doc.add_page_break()
    doc.add_heading("Cross-Feature Journeys", level=1)
    add_para(doc, "Execute these as complete product flows. A journey fails if an upstream screen succeeds but downstream state, attribution, security or notifications are wrong.")
    for case_id, title, action, expected in qa.E2E_CASES:
        doc.add_heading(f"{case_id}: {title}", level=2)
        steps = [x.strip() for x in action.split("->")]
        add_flow(doc, steps)
        add_callout(doc, "Journey pass condition", expected, fill=GREEN, accent="235C34")
        add_bullets(doc, [
            "At each arrow, verify the visible state and the exact API request/response.",
            "Verify the tenant ID and foreign keys of the newly created or updated record.",
            "Verify history, attribution, notification, email, scheduler or external provider output separately.",
            "Repeat the sensitive step with Agency B IDs and an unauthorized role.",
        ])


def impact_guide(doc):
    doc.add_page_break()
    doc.add_heading("Impact and Regression Guide", level=1)
    add_para(doc, "Use this section for change-based testing. The listed dependencies are the minimum retest surface, not an optional suggestion.")
    rows = []
    for feature in qa.FEATURES:
        rows.append((f"If {feature['name']} changes", f"Retest: primary workflow; validation; roles; tenant isolation; delete/idempotency; {feature['dependencies']}."))
    add_two_col_table(doc, rows, widths=(2.2, 4.45), header=("Change Area", "Minimum Retest Scope"))


def repro_templates(doc):
    doc.add_page_break()
    doc.add_heading("Reusable Bug Reproduction Templates", level=1)
    templates = [
        ("Permission error", "Login as a lower role -> open UI/direct endpoint -> attempt action -> verify 403/404 -> verify no DB/audit/side effect."),
        ("Validation failure", "Submit a valid baseline -> change one field to invalid/missing -> submit -> verify 400/field error -> compare DB -> verify form retains correctable values."),
        ("Duplicate submission", "Prepare valid form -> double-click or send two concurrent identical requests -> refresh -> count rows/history/notifications/provider calls."),
        ("Invalid state transition", "Create entity in known state -> call transition with skipped/terminal/reverse state -> verify current rule -> compare state/history/downstream data."),
        ("Data ownership violation", "Login Agency A -> substitute Agency B target and related IDs -> list/detail/mutate/download -> verify no data and no mutation."),
        ("Deletion behavior", "Create unreferenced and referenced records -> delete as agent/manager -> inspect CASCADE/PROTECT/SET_NULL -> repeat delete -> verify 403/204/404 and rollback."),
        ("External service failure", "Mock timeout/401/429/5xx/invalid JSON -> run action -> inspect local pending/failed state -> recover provider -> retry -> prove no duplicate remote success."),
        ("Background job failure", "Seed due valid/invalid records -> run command twice -> inspect counts/markers/errors -> simulate second worker -> prove logical idempotency."),
    ]
    for title, steps in templates:
        doc.add_heading(title, level=2)
        add_two_col_table(doc, [
            ("Feature", "Fill in."),
            ("Scenario", title),
            ("Precondition", "State, role, tenant and fixture IDs."),
            ("Steps to Reproduce", steps),
            ("Expected", "Current implementation pass condition from the relevant feature chapter/workbook row."),
            ("Actual", "Fill in with exact response/UI/database/side effects."),
            ("Evidence", "Screenshot, HTTP trace, SQL diff, log and provider mock count."),
        ], header=("Field", "Template"))


def risks_and_unknowns(doc):
    doc.add_page_break()
    doc.add_heading("Probable Risks and Clarification Register", level=1)
    add_callout(doc, "Important", "These items are not silently converted into expected behavior. QA should record the current observed result and product/engineering should decide the intended contract.", fill=AMBER, accent="8A4B08")
    doc.add_heading("Probable Architectural / QA Risks", level=2)
    add_bullets(doc, qa.PROBABLE_RISKS)
    doc.add_heading("Requires Clarification / Not Determined", level=2)
    add_bullets(doc, qa.UNKNOWNS)
    doc.add_heading("Current Implementation vs Recommended Behavior", level=2)
    add_two_col_table(doc, [
        ("Current", "Record the behavior actually produced by the reviewed code and deployed test environment."),
        ("Expected", "Use current explicit business rules and test assertions. Do not invent desired endpoints or transitions."),
        ("Recommended", "Log improvement ideas separately. Examples include optimistic locking, expiring customer sessions, durable queues, malware scanning and monitored object storage."),
        ("Bug", "A divergence from current documented/explicit behavior, data integrity, security boundary or consistent product invariant."),
        ("Clarification", "A behavior the code does not clearly define, such as whether public appointment slots are exclusive."),
    ], header=("Label", "How QA Should Use It"))


def closing(doc):
    doc.add_page_break()
    doc.add_heading("Release Readiness Checklist", level=1)
    add_bullets(doc, [
        "All P0 cases executed and passed; no open critical tenant-isolation/authentication defect.",
        "P1 failures have owner, impact, workaround and explicit release decision.",
        "Main E2E journeys pass in Agency A and tenant-isolation variants pass with Agency B.",
        "Scheduler commands processed due work once and recovered from provider/email failure.",
        "Meta and Stripe webhook signatures and replay behavior validated in test mode.",
        "Public storefront exposes only active, paid, published, fresh, eligible tenant content.",
        "Verification documents, customer records and staff-only data are absent from public payloads.",
        "Property/deal/lease journeys do not invoke payment processing; Stripe remains SaaS subscription only.",
        "Regression workbook status and bug links are complete and evidence is retained.",
        "Unknowns relevant to release environment are resolved or formally accepted.",
    ])
    add_callout(doc, "Master documentation pair", "Use Feature_Testing_Guide.docx to execute and understand the system. Use System_Test_Cases.xlsx to select cases, track results, prioritize regression and link defects.")


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    metrics = qa.summary()
    cover_page(doc, metrics)
    front_matter(doc, metrics)
    for i, feature in enumerate(qa.FEATURES, 1):
        feature_chapter(doc, i, feature)
    cross_feature_section(doc)
    impact_guide(doc)
    repro_templates(doc)
    risks_and_unknowns(doc)
    closing(doc)

    props = doc.core_properties
    props.title = "Nexora RealtyOS Feature Testing Guide"
    props.subject = "Master QA manual execution handbook"
    props.author = "Nexora RealtyOS QA"
    props.keywords = "Nexora, RealtyOS, QA, testing, Nepal real estate"
    props.comments = "Generated from current repository implementation analysis."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
